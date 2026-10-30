import csv
import io

from openpyxl import load_workbook
from fastapi.testclient import TestClient


def auth(token):
    return {"Authorization": f"Bearer {token}"}


def test_search_inventory_by_component_name(client, staff_token):
    response = client.get(
        "/api/inventory?search=Arduino Uno",
        headers=auth(staff_token),
    )

    assert response.status_code == 200

    items = response.json()

    assert len(items) == 1
    assert items[0]["name"] == "Arduino Uno"
    assert items[0]["sku"] == "EIMS-00001"


def test_search_inventory_by_sku(client, staff_token):
    response = client.get(
        "/api/inventory?search=EIMS-00002",
        headers=auth(staff_token),
    )

    assert response.status_code == 200

    items = response.json()

    assert len(items) == 1
    assert items[0]["sku"] == "EIMS-00002"
    assert items[0]["name"] == "ESP32 Dev Board"


def test_search_inventory_by_category(client, staff_token):
    response = client.get(
        "/api/inventory?search=Microcontroller",
        headers=auth(staff_token),
    )

    assert response.status_code == 200

    items = response.json()

    assert len(items) == 2
    assert all(
        item["category"] == "Microcontroller"
        for item in items
    )


def test_search_inventory_with_no_matches(client, staff_token):
    response = client.get(
        "/api/inventory?search=NonexistentComponent123",
        headers=auth(staff_token),
    )

    assert response.status_code == 200
    assert response.json() == []



def test_availability_report_classifies_all_stock_statuses(client, admin_token):
    response = client.get(
        "/api/reports/availability",
        headers=auth(admin_token),
    )

    assert response.status_code == 200

    items = {item["name"]: item for item in response.json()}

    assert items["Arduino Uno"]["status"] == "in-stock"
    assert items["ESP32 Dev Board"]["status"] == "low-stock"
    assert items["220 Ohm Resistor"]["status"] == "out-of-stock"


def test_availability_report_reflects_updated_quantity(client, admin_token):
    headers = auth(admin_token)

    # Get the current component details.
    response = client.get("/api/reports/availability", headers=headers)
    assert response.status_code == 200

    arduino = next(
        item for item in response.json()
        if item["name"] == "Arduino Uno"
    )

    # Change its quantity to zero.
    update = client.post(
    f"/api/inventory/{arduino['id']}/movements",
    json={
        "movement_type": "STOCK_OUT",
        "quantity": arduino["quantity"],
        "reason": "Testing availability report after stock depletion",
    },
    headers=headers,
)
    assert update.status_code == 201

    # The report should show the updated quantity and status.
    report = client.get("/api/reports/availability", headers=headers)
    updated_arduino = next(
        item for item in report.json()
        if item["name"] == "Arduino Uno"
    )

    assert updated_arduino["quantity"] == 0
    assert updated_arduino["status"] == "out-of-stock"


def test_availability_report_excludes_inactive_components(client, admin_token):
    headers = auth(admin_token)

    response = client.get("/api/reports/availability", headers=headers)
    arduino = next(
        item for item in response.json()
        if item["name"] == "Arduino Uno"
    )

    # Delete the component (the API marks it inactive).
    deleted = client.delete(
        f"/api/inventory/{arduino['id']}",
        headers=headers,
    )
    assert deleted.status_code == 200

    # Inactive components should not appear in the report.
    report = client.get("/api/reports/availability", headers=headers)
    names = [item["name"] for item in report.json()]

    assert "Arduino Uno" not in names


def test_availability_report_filters_in_stock(client, admin_token):
    response = client.get(
        "/api/reports/availability?status=in-stock",
        headers=auth(admin_token),
    )

    assert response.status_code == 200
    assert response.json()
    assert all(item["status"] == "in-stock" for item in response.json())


def test_availability_report_filters_low_stock(client, admin_token):
    response = client.get(
        "/api/reports/availability?status=low-stock",
        headers=auth(admin_token),
    )

    assert response.status_code == 200
    assert response.json()
    assert all(item["status"] == "low-stock" for item in response.json())


def test_availability_report_filters_out_of_stock(client, admin_token):
    response = client.get(
        "/api/reports/availability?status=out-of-stock",
        headers=auth(admin_token),
    )

    assert response.status_code == 200
    assert response.json()
    assert all(item["status"] == "out-of-stock" for item in response.json())


def test_availability_report_rejects_invalid_status(client, admin_token):
    response = client.get(
        "/api/reports/availability?status=invalid",
        headers=auth(admin_token),
    )

    assert response.status_code == 400



def test_availability_report_returns_matching_results(client, admin_token):
    response = client.get(
        "/api/reports/availability?status=low-stock",
        headers=auth(admin_token),
    )

    assert response.status_code == 200
    assert response.json()
    assert all(item["status"] == "low-stock" for item in response.json())

def test_export_availability_report_as_csv(client, admin_token):
    response = client.get(
        "/api/reports/availability/export?format=csv",
        headers=auth(admin_token),
    )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/csv")
    assert "attachment" in response.headers["content-disposition"]

    rows = list(csv.reader(io.StringIO(response.text)))

    assert rows[0] == [
        "SKU", "Name", "Category", "Quantity",
        "Unit Price", "Supplier", "Reorder Level", "Stock Status"
    ]
    assert len(rows) > 1
    assert any(row[1] == "Arduino Uno" for row in rows[1:])


def test_export_availability_report_as_excel(client, admin_token):
    response = client.get(
        "/api/reports/availability/export?format=excel",
        headers=auth(admin_token),
    )

    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]

    workbook = load_workbook(io.BytesIO(response.content))
    worksheet = workbook.active

    assert worksheet["A1"].value == "SKU"
    assert worksheet["B1"].value == "Name"
    assert worksheet.max_row > 1
    assert any(
        row[1] == "Arduino Uno"
        for row in worksheet.iter_rows(min_row=2, values_only=True)
    )


def test_export_availability_report_as_pdf(client, admin_token):
    response = client.get(
        "/api/reports/availability/export?format=pdf",
        headers=auth(admin_token),
    )

    assert response.status_code == 200
    assert response.content.startswith(b"%PDF-")
    assert response.headers["content-type"] == "application/pdf"
    assert "attachment" in response.headers["content-disposition"]


def test_export_availability_report_with_stock_filter(client, admin_token):
    response = client.get(
        "/api/reports/availability/export?format=csv&status=low-stock",
        headers=auth(admin_token),
    )

    assert response.status_code == 200

    rows = list(csv.DictReader(io.StringIO(response.text)))

    assert rows
    assert all(row["Stock Status"] == "low-stock" for row in rows)


def test_export_availability_report_rejects_invalid_format(client, admin_token):
    response = client.get(
        "/api/reports/availability/export?format=doc",
        headers=auth(admin_token),
    )

    assert response.status_code == 422


def test_staff_cannot_access_availability_report_or_exports(client, staff_token):
    report = client.get(
        "/api/reports/availability",
        headers=auth(staff_token),
    )
    export = client.get(
        "/api/reports/availability/export?format=csv",
        headers=auth(staff_token),
    )

    assert report.status_code == 403
    assert export.status_code == 403